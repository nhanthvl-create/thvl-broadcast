from flask import Flask, render_template, request, redirect, url_for, flash, session, jsonify, send_file
from flask_sqlalchemy import SQLAlchemy
from werkzeug.security import generate_password_hash, check_password_hash
from datetime import datetime, date, timedelta
from functools import wraps
import os, re, io, shutil, threading, time as ttime

import os as _os
_base_dir = _os.path.dirname(_os.path.abspath(__file__))
app = Flask(__name__,
            template_folder=_os.path.join(_base_dir, 'templates'),
            static_folder=_os.path.join(_base_dir, 'static') if _os.path.exists(_os.path.join(_base_dir, 'static')) else None)
app.secret_key = os.environ.get('SECRET_KEY', 'thvl-broadcast-2025-changeme')
# Hien thi loi chi tiet tren browser khi chay cloud (debug tam thoi)
if os.environ.get('RENDER'):
    app.config['PROPAGATE_EXCEPTIONS'] = False
    import logging
    logging.basicConfig(level=logging.DEBUG)

# Database: PostgreSQL neu co DATABASE_URL (Render cloud), SQLite khi chay local
_db_url = os.environ.get('DATABASE_URL', '')
if _db_url.startswith('postgres://'):
    # Render dung postgres:// nhung SQLAlchemy can postgresql://
    _db_url = _db_url.replace('postgres://', 'postgresql://', 1)
if _db_url:
    app.config['SQLALCHEMY_DATABASE_URI'] = _db_url
else:
    # Local Windows: dung SQLite trong thu muc instance/
    _sqlite_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'instance')
    os.makedirs(_sqlite_dir, exist_ok=True)
    app.config['SQLALCHEMY_DATABASE_URI'] = f'sqlite:///{_sqlite_dir}/thvl.db'
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False
app.config['MAX_CONTENT_LENGTH'] = 10 * 1024 * 1024

db = SQLAlchemy(app)

@app.before_request
def ensure_db():
    global _db_initialized
    if not _db_initialized:
        try:
            init_db()
            _db_initialized = True
        except Exception as e:
            print(f'DB init error: {e}')

@app.route('/health')
def health():
    try:
        db.session.execute(db.text('SELECT 1'))
        return 'OK - DB connected', 200
    except Exception as e:
        return f'DB ERROR: {e}', 500

@app.errorhandler(500)
def handle_500(e):
    import traceback
    tb = traceback.format_exc()
    print('500 ERROR:', tb)
    return f'<pre>500 Error:\n{tb}</pre>', 500

# ═══════════════════════════════════════════════════
#  MODELS
# ═══════════════════════════════════════════════════

class User(db.Model):
    id            = db.Column(db.Integer, primary_key=True)
    username      = db.Column(db.String(50), unique=True, nullable=False)
    full_name     = db.Column(db.String(100), nullable=False)
    password_hash = db.Column(db.String(200), nullable=False)
    phong_ban     = db.Column(db.String(100), nullable=False)
    role          = db.Column(db.String(20), nullable=False)
    # roles: admin | phongct | phongcm | phatsong
    active        = db.Column(db.Boolean, default=True)
    created_at    = db.Column(db.DateTime, default=datetime.utcnow)

    def set_password(self, pw): self.password_hash = generate_password_hash(pw)
    def check_password(self, pw): return check_password_hash(self.password_hash, pw)
    def initials(self): return self.full_name[:2].upper()


class LichHDPS(db.Model):
    __tablename__ = 'lich_hdps'
    id          = db.Column(db.Integer, primary_key=True)
    ngay_phat   = db.Column(db.Date, nullable=False)
    kenh        = db.Column(db.String(10), default='THVL1')
    ghi_chu     = db.Column(db.Text)
    trang_thai  = db.Column(db.String(20), default='nhap')  # nhap | hoan_chinh | da_xuat
    nguoi_tao_id= db.Column(db.Integer, db.ForeignKey('user.id'))
    nguoi_tao   = db.relationship('User')
    created_at  = db.Column(db.DateTime, default=datetime.utcnow)
    updated_at  = db.Column(db.DateTime, default=datetime.utcnow, onupdate=datetime.utcnow)
    items       = db.relationship('LichItem', backref='lich', cascade='all,delete-orphan',
                                  order_by='LichItem.stt')

    __table_args__ = (db.UniqueConstraint('ngay_phat', 'kenh', name='uq_ngay_kenh'),)


class LichItem(db.Model):
    """Một dòng trong lịch HDPS."""
    id          = db.Column(db.Integer, primary_key=True)
    lich_id     = db.Column(db.Integer, db.ForeignKey('lich_hdps.id'), nullable=False)
    stt         = db.Column(db.Integer, nullable=False)
    gio_phat    = db.Column(db.String(12))          # HH:MM:SS computed
    tl_ct       = db.Column(db.String(12))          # thời lượng CT  HH:MM:SS
    tl_qc       = db.Column(db.String(12))          # thời lượng QC  HH:MM:SS (có thể trống)
    ten_ct      = db.Column(db.String(300), nullable=False)
    loai        = db.Column(db.String(20), default='file')
    # loai: file | truc_tiep | phat_lai | phong_cm
    duong_dan   = db.Column(db.Text)                # path server
    ten_file    = db.Column(db.Text)                # tên file đầy đủ (nhiều file, xuống dòng)
    tai_tro     = db.Column(db.String(20), default='khong')  # co | khong
    trang_thai  = db.Column(db.String(20), default='cho')
    # cho | co_file | da_xac_nhan | thieu_file
    phieu_nt_id     = db.Column(db.Integer, db.ForeignKey('phieu_nghiem_thu.id'), nullable=True)
    phieu_nt        = db.relationship('PhieuNghiemThu')
    ghi_chu         = db.Column(db.Text)
    nguoi_thao_tac  = db.Column(db.String(100))   # Người gán file vào lịch
    file_size_mb    = db.Column(db.Float)          # Dung lượng file (MB)
    thoi_luong_giay = db.Column(db.Integer)        # Thời lượng thực tế từ file (giây)


class PhieuNghiemThu(db.Model):
    id            = db.Column(db.Integer, primary_key=True)
    ten_ct        = db.Column(db.String(200), nullable=False)
    phong_ban     = db.Column(db.String(100), nullable=False)
    ten_file      = db.Column(db.Text, nullable=False)
    duong_dan     = db.Column(db.Text, nullable=False)
    ngay_phat     = db.Column(db.Date, nullable=False)
    thoi_luong    = db.Column(db.String(20))
    tai_tro       = db.Column(db.String(20), default='Không tài trợ')
    ghi_chu       = db.Column(db.Text)
    trang_thai    = db.Column(db.String(20), default='cho_duyet')
    # cho_duyet | da_duyet | tra_ve
    ly_do_tra     = db.Column(db.Text)
    nguoi_nop_id  = db.Column(db.Integer, db.ForeignKey('user.id'))
    nguoi_nop     = db.relationship('User', foreign_keys=[nguoi_nop_id])
    nguoi_duyet_id= db.Column(db.Integer, db.ForeignKey('user.id'))
    nguoi_duyet   = db.relationship('User', foreign_keys=[nguoi_duyet_id])
    created_at    = db.Column(db.DateTime, default=datetime.utcnow)
    updated_at    = db.Column(db.DateTime, default=datetime.utcnow, onupdate=datetime.utcnow)


class CopyJob(db.Model):
    """Theo dõi việc copy file về máy phát sóng."""
    id           = db.Column(db.Integer, primary_key=True)
    lich_id      = db.Column(db.Integer, db.ForeignKey('lich_hdps.id'))
    lich         = db.relationship('LichHDPS')
    dest_path    = db.Column(db.String(300))         # đường dẫn đích trên máy phát sóng
    trang_thai   = db.Column(db.String(20), default='cho')
    # cho | dang_chay | hoan_thanh | loi
    tong_file    = db.Column(db.Integer, default=0)
    da_copy      = db.Column(db.Integer, default=0)
    log          = db.Column(db.Text)
    nguoi_tao_id = db.Column(db.Integer, db.ForeignKey('user.id'))
    created_at   = db.Column(db.DateTime, default=datetime.utcnow)
    updated_at   = db.Column(db.DateTime, default=datetime.utcnow, onupdate=datetime.utcnow)

class SysConfig(db.Model):
    key   = db.Column(db.String(50), primary_key=True)
    value = db.Column(db.Text)

# ═══════════════════════════════════════════════════
#  HELPERS
# ═══════════════════════════════════════════════════

def get_cfg(key, default=''):
    c = SysConfig.query.get(key)
    return c.value if c else default

def set_cfg(key, val):
    c = SysConfig.query.get(key)
    if c: c.value = val
    else: db.session.add(SysConfig(key=key, value=val))
    db.session.commit()

def tc_to_sec(tc):
    """'HH:MM:SS' → seconds float."""
    if not tc: return 0
    parts = tc.strip().split(':')
    try:
        if len(parts) == 3: return int(parts[0])*3600 + int(parts[1])*60 + float(parts[2])
        if len(parts) == 2: return int(parts[0])*60 + float(parts[1])
        return float(parts[0])
    except: return 0

def sec_to_tc(s):
    """seconds → 'HH:MM:SS'."""
    s = int(s)
    return f"{s//3600:02d}:{(s%3600)//60:02d}:{s%60:02d}"

def add_tc(base_tc, dur_tc, qc_tc=''):
    """Add durations to a base timecode."""
    total = tc_to_sec(base_tc) + tc_to_sec(dur_tc) + tc_to_sec(qc_tc)
    return sec_to_tc(total)

def recalc_times(lich_id):
    """Recalculate gio_phat for all items in a lich."""
    items = LichItem.query.filter_by(lich_id=lich_id).order_by(LichItem.stt).all()
    if not items: return
    cursor = tc_to_sec(items[0].gio_phat or '05:00:00')
    for i, item in enumerate(items):
        if i == 0:
            pass  # keep first item's time
        else:
            item.gio_phat = sec_to_tc(cursor)
        cursor += tc_to_sec(item.tl_ct) + tc_to_sec(item.tl_qc)
    db.session.commit()

def parse_excel_hdps(filepath, ngay_phat, kenh, user_id):
    """Parse HDPS Excel file using openpyxl/xlrd only (no pandas)."""
    import os

    # Try openpyxl first (.xlsx), fallback to xlrd (.xls)
    rows = []
    ext = os.path.splitext(filepath)[1].lower()

    if ext == '.xlsx':
        # Dung openpyxl cho .xlsx
        import openpyxl
        from openpyxl import load_workbook
        wb = load_workbook(filepath, data_only=True)
        ws = wb.active
        for row in ws.iter_rows(values_only=True):
            rows.append(list(row))
    else:
        # Dung xlrd cho .xls (dinh dang cu)
        import xlrd
        try:
            wb = xlrd.open_workbook(filepath)
        except Exception:
            # Neu .xls that ra la .xlsx doi ten va thu lai
            import shutil
            new_path = filepath + 'x'
            shutil.copy(filepath, new_path)
            import openpyxl
            wb2 = openpyxl.load_workbook(new_path, data_only=True)
            ws2 = wb2.active
            for row in ws2.iter_rows(values_only=True):
                rows.append(list(row))
            os.remove(new_path)
            # skip xlrd section
            wb = None
        if wb is not None:
            ws = wb.sheet_by_index(0)
            for i in range(ws.nrows):
                row = []
                for j in range(ws.ncols):
                    cell = ws.cell(i, j)
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        try:
                            t = xlrd.xldate_as_tuple(cell.value, wb.datemode)
                            class _T:
                                def __init__(s,h,m,sc): s.hour=h; s.minute=m; s.second=int(sc)
                            row.append(_T(t[3], t[4], t[5]))
                        except Exception:
                            # Fallback: parse as fraction of day
                            frac = cell.value % 1
                            total_sec = int(frac * 86400)
                            class _T2:
                                def __init__(s,h,m,sc): s.hour=h; s.minute=m; s.second=sc
                            row.append(_T2(total_sec//3600, (total_sec%3600)//60, total_sec%60))
                    else:
                        row.append(cell.value)
                rows.append(row)

    # Delete existing lich if any
    lich = LichHDPS.query.filter_by(ngay_phat=ngay_phat, kenh=kenh).first()
    if lich:
        LichItem.query.filter_by(lich_id=lich.id).delete()
    else:
        lich = LichHDPS(ngay_phat=ngay_phat, kenh=kenh,
                        trang_thai='nhap', nguoi_tao_id=user_id)
        db.session.add(lich)
        db.session.flush()

    items_added = 0
    cursor_sec  = 5 * 3600  # default 05:00:00
    stt = 1

    def _tc_val(v):
        """Extract seconds from a cell value (time object or string HH:MM:SS)."""
        if v is None: return 0
        if hasattr(v, 'hour'):
            return v.hour * 3600 + v.minute * 60 + getattr(v, 'second', 0)
        s = str(v).strip()
        if not s or s == 'None': return 0
        parts = s.split(':')
        try:
            if len(parts) == 3: return int(parts[0])*3600 + int(parts[1])*60 + int(float(parts[2]))
            if len(parts) == 2: return int(parts[0])*60 + int(parts[1])
        except: pass
        return 0

    for row in rows:
        # Pad row to at least 6 cols
        while len(row) < 6: row.append(None)

        ten_ct_raw = str(row[3]).strip() if row[3] is not None else ''

        # Skip empty / header / footer rows
        if not ten_ct_raw or ten_ct_raw in ('None','nan'):
            continue
        skip_words = ['BM0','BM03','Người lên lịch','Lãnh đạo','TLQC','Tổng','GIỜ PHÁT','GIO PHAT','STT']
        if any(ten_ct_raw.upper().startswith(w.upper()) for w in skip_words):
            continue

        # Parse gio_phat
        v0 = row[0]
        if hasattr(v0, 'hour'):
            cursor_sec = v0.hour*3600 + v0.minute*60 + getattr(v0,'second',0)

        gio_str = sec_to_tc(cursor_sec)

        tl_ct_sec = _tc_val(row[1])
        tl_ct_str = sec_to_tc(tl_ct_sec)
        tl_qc_sec = _tc_val(row[2])
        tl_qc_str = sec_to_tc(tl_qc_sec) if tl_qc_sec else ''

        duong_dan_raw = str(row[4]).strip() if row[4] is not None else ''
        if duong_dan_raw == 'None': duong_dan_raw = ''

        # Detect loai
        dd_lower = duong_dan_raw.lower()
        ten_lower = ten_ct_raw.lower()
        if 'trực tiếp' in dd_lower or 'truc tiep' in dd_lower:
            loai = 'truc_tiep'
        elif 'phát lại' in dd_lower or 'phat lai' in dd_lower or 'phát lại' in ten_lower:
            loai = 'phat_lai'
        elif duong_dan_raw.upper() in ('P.CM','P.TS','P.SXCT','P.BD','P.CT','P.CM','PCM'):
            loai = 'phong_cm'
        else:
            loai = 'file'

        trang_thai = 'truc_tiep' if loai == 'truc_tiep' else 'cho'

        item = LichItem(
            lich_id    = lich.id,
            stt        = stt,
            gio_phat   = gio_str,
            tl_ct      = tl_ct_str,
            tl_qc      = tl_qc_str,
            ten_ct     = ten_ct_raw,
            loai       = loai,
            duong_dan  = duong_dan_raw,
            tai_tro    = 'co' if 'tài trợ' in ten_lower and 'không' not in ten_lower else 'khong',
            trang_thai = trang_thai,
        )
        db.session.add(item)
        stt += 1
        items_added += 1
        cursor_sec += tl_ct_sec + tl_qc_sec

    db.session.commit()
    return lich, items_added
def gen_ply(lich, dest_path):
    """Generate PlayBox AirBox .ply playlist content from LichHDPS."""
    import time as ttime
    lines = []
    ts = str(int(ttime.time() * 1000))

    lines.append(f'#FILENAME {dest_path}\\{lich.ngay_phat.strftime("%d%m%Y")}_{lich.kenh}.ply')
    lines.append(f'#PLAYLIST_FILE_NAME {dest_path}\\{lich.ngay_phat.strftime("%d%m%Y")}_{lich.kenh}.ply')
    lines.append(f'#PLAYLISTID {ts[:12]}')
    lines.append(f'#PLAYLISTTC 00:00:00:00')

    items = LichItem.query.filter_by(lich_id=lich.id).order_by(LichItem.stt).all()

    for idx, item in enumerate(items):
        # STARTTIME for first item or items with hard start
        if idx == 0 or item.loai == 'stop':
            gio = item.gio_phat or '05:00:00'
            lines.append(f'#STARTTIME 0;{gio};-1;-1;0;;0')

        if item.loai == 'truc_tiep':
            lines.append(f'#LISTID {ts}{idx:04d}')
            lines.append('#DYNAMICMEDIA FALSE')
            lines.append('#EVENT STOP')
            continue

        if item.loai == 'stop':
            lines.append(f'#LISTID {ts}{idx:04d}')
            lines.append('#DYNAMICMEDIA FALSE')
            lines.append('#EVENT STOP')
            continue

        if not item.ten_file:
            continue

        file_lines = [f.strip() for f in item.ten_file.strip().split('\n') if f.strip()]
        base_path  = (item.duong_dan or '').rstrip('\\')

        for fi, fname in enumerate(file_lines):
            # Ensure .mpg extension
            if not fname.lower().endswith('.mpg'):
                fname = fname + '.mpg'
            full_path = f'{base_path}\\{fname}'
            listid = f'{ts}{idx:04d}{fi}'
            lines.append(f'#LISTID {listid}')
            lines.append('#DYNAMICMEDIA FALSE')
            lines.append('#TYPE Clip')
            lines.append('#TC 0.00000')
            display = fname.replace('.mpg', '')
            lines.append(f'"{full_path}";0.00000;0.00000;;{display}')

    return '\r\n'.join(lines) + '\r\n'


# ═══════════════════════════════════════════════════
#  AUTH DECORATORS
# ═══════════════════════════════════════════════════

def login_required(f):
    @wraps(f)
    def dec(*a, **kw):
        if 'user_id' not in session:
            return redirect(url_for('login'))
        return f(*a, **kw)
    return dec

def roles(*allowed):
    def deco(f):
        @wraps(f)
        def dec(*a, **kw):
            if session.get('role') not in allowed:
                flash('Bạn không có quyền thực hiện thao tác này.', 'danger')
                return redirect(url_for('dashboard'))
            return f(*a, **kw)
        return dec
    return deco

app.jinja_env.globals.update(get_current_user=lambda: User.query.get(session.get('user_id')),
                              now=datetime.now, tc_to_sec=tc_to_sec)

# ═══════════════════════════════════════════════════
#  AUTH ROUTES
# ═══════════════════════════════════════════════════

@app.route('/')
def index():
    return redirect(url_for('dashboard') if 'user_id' in session else url_for('login'))

@app.route('/login', methods=['GET','POST'])
def login():
    if request.method == 'POST':
        u = User.query.filter_by(username=request.form['username'].strip(), active=True).first()
        if u and u.check_password(request.form['password']):
            session.update(user_id=u.id, username=u.username, full_name=u.full_name,
                           role=u.role, phong_ban=u.phong_ban)
            return redirect(url_for('dashboard'))
        flash('Tên đăng nhập hoặc mật khẩu không đúng.', 'danger')
    return render_template('login.html')

@app.route('/logout')
def logout():
    session.clear(); return redirect(url_for('login'))

# ═══════════════════════════════════════════════════
#  DASHBOARD
# ═══════════════════════════════════════════════════

@app.route('/dashboard')
@login_required
def dashboard():
    today = date.today()
    # Lịch hôm nay + ngày mai
    lich_today = LichHDPS.query.filter_by(ngay_phat=today).all()
    lich_tom   = LichHDPS.query.filter_by(ngay_phat=today + timedelta(days=1)).all()
    # Phiếu NT chờ duyệt
    cho_duyet  = PhieuNghiemThu.query.filter_by(trang_thai='cho_duyet').count()
    # Items thiếu file (hôm nay)
    thieu_file = 0
    for l in lich_today:
        thieu_file += LichItem.query.filter_by(lich_id=l.id, trang_thai='thieu_file').count()
        thieu_file += LichItem.query.filter_by(lich_id=l.id, trang_thai='cho').count()
    recent_phieu = PhieuNghiemThu.query.order_by(
        PhieuNghiemThu.updated_at.desc()).limit(10).all()
    copy_jobs = CopyJob.query.order_by(CopyJob.created_at.desc()).limit(5).all()
    return render_template('dashboard.html', today=today,
        lich_today=lich_today, lich_tom=lich_tom,
        cho_duyet=cho_duyet, thieu_file=thieu_file,
        recent_phieu=recent_phieu, copy_jobs=copy_jobs)

# ═══════════════════════════════════════════════════
#  LỊCH HDPS
# ═══════════════════════════════════════════════════

@app.route('/lich')
@login_required
def lich_list():
    q = LichHDPS.query
    kenh = request.args.get('kenh')
    if kenh: q = q.filter_by(kenh=kenh)
    lichs = q.order_by(LichHDPS.ngay_phat.desc()).limit(30).all()
    return render_template('lich_list.html', lichs=lichs)

@app.route('/lich/new', methods=['GET','POST'])
@login_required
@roles('admin','phongct')
def lich_new():
    if request.method == 'POST':
        ngay = datetime.strptime(request.form['ngay_phat'], '%Y-%m-%d').date()
        kenh = request.form.get('kenh','THVL1')
        existing = LichHDPS.query.filter_by(ngay_phat=ngay, kenh=kenh).first()
        if existing:
            flash(f'Đã có lịch {kenh} ngày {ngay.strftime("%d/%m/%Y")}.', 'warning')
            return redirect(url_for('lich_detail', lid=existing.id))
        lich = LichHDPS(ngay_phat=ngay, kenh=kenh,
                        ghi_chu=request.form.get('ghi_chu',''),
                        trang_thai='nhap', nguoi_tao_id=session['user_id'])
        db.session.add(lich)
        db.session.commit()
        flash('Đã tạo lịch mới.', 'success')
        return redirect(url_for('lich_detail', lid=lich.id))
    return render_template('lich_new.html')

@app.route('/lich/upload', methods=['GET','POST'])
@login_required
@roles('admin','phongct')
def lich_upload():
    if request.method == 'POST':
        f = request.files.get('excel_file')
        if not f or not f.filename.endswith(('.xls','.xlsx')):
            flash('Vui lòng chọn file Excel (.xls hoặc .xlsx).', 'danger')
            return redirect(request.url)
        ngay = datetime.strptime(request.form['ngay_phat'], '%Y-%m-%d').date()
        kenh = request.form.get('kenh', 'THVL1')
        tmp = f'/tmp/hdps_upload_{session["user_id"]}.xls'
        f.save(tmp)
        try:
            lich, n = parse_excel_hdps(tmp, ngay, kenh, session['user_id'])
            flash(f'Import thành công {n} mục vào lịch {kenh} {ngay.strftime("%d/%m/%Y")}.', 'success')
            return redirect(url_for('lich_detail', lid=lich.id))
        except Exception as e:
            flash(f'Lỗi đọc file Excel: {e}', 'danger')
            return redirect(request.url)
        finally:
            if os.path.exists(tmp): os.remove(tmp)
    return render_template('lich_upload.html')

@app.route('/lich/<int:lid>')
@login_required
def lich_detail(lid):
    lich  = LichHDPS.query.get_or_404(lid)
    items = LichItem.query.filter_by(lich_id=lid).order_by(LichItem.stt).all()
    # Thống kê
    stats = {
        'tong': len(items),
        'co_file': sum(1 for i in items if i.trang_thai == 'co_file'),
        'da_xac_nhan': sum(1 for i in items if i.trang_thai == 'da_xac_nhan'),
        'cho': sum(1 for i in items if i.trang_thai == 'cho'),
        'thieu': sum(1 for i in items if i.trang_thai == 'thieu_file'),
        'truc_tiep': sum(1 for i in items if i.loai == 'truc_tiep'),
    }
    tl_qc_total = sum(tc_to_sec(i.tl_qc) for i in items)
    copy_jobs = CopyJob.query.filter_by(lich_id=lid).order_by(CopyJob.created_at.desc()).all()
    all_users = User.query.filter_by(active=True).order_by(User.full_name).all()
    return render_template('lich_detail.html', lich=lich, items=items,
                           stats=stats, tl_qc_total=sec_to_tc(tl_qc_total),
                           copy_jobs=copy_jobs, all_users=all_users)

@app.route('/lich/<int:lid>/item/add', methods=['POST'])
@login_required
@roles('admin','phongct')
def item_add(lid):
    lich  = LichHDPS.query.get_or_404(lid)
    max_stt = db.session.query(db.func.max(LichItem.stt)).filter_by(lich_id=lid).scalar() or 0
    item = LichItem(
        lich_id   = lid,
        stt       = max_stt + 1,
        gio_phat  = request.form.get('gio_phat',''),
        tl_ct     = request.form.get('tl_ct',''),
        tl_qc     = request.form.get('tl_qc',''),
        ten_ct    = request.form.get('ten_ct','').strip(),
        loai      = request.form.get('loai','file'),
        duong_dan = request.form.get('duong_dan','').strip(),
        ten_file  = request.form.get('ten_file','').strip(),
        tai_tro   = request.form.get('tai_tro','khong'),
        ghi_chu   = request.form.get('ghi_chu','').strip(),
        trang_thai = 'truc_tiep' if request.form.get('loai')=='truc_tiep' else 'cho',
    )
    db.session.add(item)
    db.session.commit()
    flash('Đã thêm mục vào lịch.', 'success')
    return redirect(url_for('lich_detail', lid=lid))

@app.route('/lich/item/<int:iid>/edit', methods=['GET','POST'])
@login_required
@roles('admin','phongct')
def item_edit(iid):
    item = LichItem.query.get_or_404(iid)
    if request.method == 'POST':
        item.gio_phat  = request.form.get('gio_phat', item.gio_phat)
        item.tl_ct     = request.form.get('tl_ct', item.tl_ct)
        item.tl_qc     = request.form.get('tl_qc', item.tl_qc)
        item.ten_ct    = request.form.get('ten_ct','').strip()
        item.loai      = request.form.get('loai', item.loai)
        item.duong_dan = request.form.get('duong_dan','').strip()
        item.ten_file  = request.form.get('ten_file','').strip()
        item.tai_tro   = request.form.get('tai_tro','khong')
        item.ghi_chu   = request.form.get('ghi_chu','').strip()
        if item.loai == 'truc_tiep': item.trang_thai = 'truc_tiep'
        db.session.commit()
        flash('Đã cập nhật.', 'success')
        return redirect(url_for('lich_detail', lid=item.lich_id))
    return render_template('item_edit.html', item=item)

@app.route('/lich/item/<int:iid>/xac_nhan', methods=['POST'])
@login_required
@roles('admin','phongct')
def item_xac_nhan(iid):
    item = LichItem.query.get_or_404(iid)
    item.trang_thai = 'da_xac_nhan'
    # Nếu có phiếu NT liên kết thì cập nhật tên file / path
    pid = request.form.get('phieu_id')
    if pid:
        p = PhieuNghiemThu.query.get(int(pid))
        if p:
            item.phieu_nt_id = p.id
            item.ten_file    = p.ten_file
            item.duong_dan   = p.duong_dan
    db.session.commit()
    return jsonify(ok=True)

@app.route('/lich/item/<int:iid>/upload_mpg', methods=['POST'])
@login_required
def item_upload_mpg(iid):
    """Upload file .mpg, doc ten + thoi luong, gan vao LichItem."""
    item = LichItem.query.get_or_404(iid)
    f = request.files.get('mpg_file')
    nguoi_tt = request.form.get('nguoi_thao_tac', '').strip()

    if not f or not f.filename:
        return jsonify(ok=False, error='Chưa chọn file'), 400

    fname = f.filename.strip()
    # Strip path (Windows may send full path)
    import ntpath
    fname = ntpath.basename(fname)

    if not fname.lower().endswith('.mpg'):
        return jsonify(ok=False, error='Chỉ chấp nhận file .mpg'), 400

    # Save temporarily to read duration
    import tempfile, subprocess
    tmp_path = None
    duration_sec = None
    file_size_mb = None

    try:
        with tempfile.NamedTemporaryFile(suffix='.mpg', delete=False) as tmp:
            f.save(tmp)
            tmp_path = tmp.name

        file_size_mb = round(os.path.getsize(tmp_path) / (1024*1024), 1)

        # Try ffprobe to get duration
        try:
            result = subprocess.run(
                ['ffprobe', '-v', 'quiet', '-print_format', 'json',
                 '-show_format', tmp_path],
                capture_output=True, text=True, timeout=15
            )
            if result.returncode == 0:
                import json
                info = json.loads(result.stdout)
                duration_sec = int(float(info.get('format', {}).get('duration', 0)))
        except Exception:
            duration_sec = None
    except Exception as e:
        return jsonify(ok=False, error=str(e)), 500
    finally:
        if tmp_path and os.path.exists(tmp_path):
            os.unlink(tmp_path)

    # Update item
    ten_file_clean = fname[:-4] if fname.lower().endswith('.mpg') else fname
    item.ten_file          = ten_file_clean
    item.nguoi_thao_tac    = nguoi_tt or session.get('full_name', '')
    item.file_size_mb      = file_size_mb
    if duration_sec:
        item.thoi_luong_giay = duration_sec
        item.tl_ct = sec_to_tc(duration_sec)
    item.trang_thai = 'co_file'
    db.session.commit()

    return jsonify(
        ok=True,
        ten_file=ten_file_clean,
        tl_ct=item.tl_ct or '',
        file_size_mb=file_size_mb,
        duration_sec=duration_sec,
        nguoi_thao_tac=item.nguoi_thao_tac,
        trang_thai=item.trang_thai
    )

@app.route('/api/users_list')
@login_required
def api_users_list():
    """Danh sach nhan vien cho dropdown nguoi thao tac."""
    users = User.query.filter_by(active=True).order_by(User.full_name).all()
    return jsonify([{'id': u.id, 'name': u.full_name, 'phong': u.phong_ban} for u in users])

@app.route('/lich/item/<int:iid>/set_operator', methods=['POST'])
@login_required
def item_set_operator(iid):
    item = LichItem.query.get_or_404(iid)
    item.nguoi_thao_tac = request.form.get('nguoi_thao_tac', '').strip()
    db.session.commit()
    return jsonify(ok=True)

@app.route('/lich/item/<int:iid>/delete', methods=['POST'])
@login_required
@roles('admin','phongct')
def item_delete(iid):
    item = LichItem.query.get_or_404(iid)
    lid  = item.lich_id
    db.session.delete(item)
    db.session.commit()
    return redirect(url_for('lich_detail', lid=lid))

@app.route('/lich/<int:lid>/hoan_chinh', methods=['POST'])
@login_required
@roles('admin','phongct')
def lich_hoan_chinh(lid):
    lich = LichHDPS.query.get_or_404(lid)
    lich.trang_thai = 'hoan_chinh'
    db.session.commit()
    flash('Lịch đã được đánh dấu hoàn chỉnh.', 'success')
    return redirect(url_for('lich_detail', lid=lid))

# ── Xuất .ply ──

@app.route('/lich/<int:lid>/xuat_ply')
@login_required
@roles('admin','phongct','phatsong')
def xuat_ply(lid):
    lich = LichHDPS.query.get_or_404(lid)
    dest = request.args.get('dest', get_cfg('dest_default', r'V:\Thanhpham_PS_K1'))
    content = gen_ply(lich, dest)
    fname = f'{lich.ngay_phat.strftime("%d%m%Y")}_{lich.kenh}.ply'
    lich.trang_thai = 'da_xuat'
    db.session.commit()
    return send_file(
        io.BytesIO(content.encode('utf-8')),
        mimetype='text/plain',
        as_attachment=True,
        download_name=fname
    )

# ── Realtime status API ──

@app.route('/api/lich/<int:lid>/status')
@login_required
def api_lich_status(lid):
    items = LichItem.query.filter_by(lich_id=lid).order_by(LichItem.stt).all()
    return jsonify([{
        'id': i.id, 'stt': i.stt, 'ten_ct': i.ten_ct,
        'gio_phat': i.gio_phat, 'loai': i.loai,
        'trang_thai': i.trang_thai, 'tai_tro': i.tai_tro,
        'tl_ct': i.tl_ct, 'tl_qc': i.tl_qc,
        'ten_file': i.ten_file or '',
        'nguoi_thao_tac': i.nguoi_thao_tac or '',
        'file_size_mb': i.file_size_mb,
    } for i in items])

# ═══════════════════════════════════════════════════
#  PHIẾU NGHIỆM THU
# ═══════════════════════════════════════════════════

@app.route('/phieu')
@login_required
def phieu_list():
    q = PhieuNghiemThu.query
    if session['role'] == 'phongcm':
        q = q.filter_by(nguoi_nop_id=session['user_id'])
    ngay = request.args.get('ngay')
    phong = request.args.get('phong')
    tt    = request.args.get('tt')
    if ngay:  q = q.filter(PhieuNghiemThu.ngay_phat == datetime.strptime(ngay,'%Y-%m-%d').date())
    if phong: q = q.filter_by(phong_ban=phong)
    if tt:    q = q.filter_by(trang_thai=tt)
    phieu_list = q.order_by(PhieuNghiemThu.ngay_phat.desc(), PhieuNghiemThu.created_at.desc()).all()
    phong_list = [p[0] for p in db.session.query(PhieuNghiemThu.phong_ban).distinct().all()]
    return render_template('phieu_list.html', phieu_list=phieu_list, phong_list=phong_list)

@app.route('/phieu/new', methods=['GET','POST'])
@login_required
@roles('admin','phongcm','phongct')
def phieu_new():
    if request.method == 'POST':
        p = PhieuNghiemThu(
            ten_ct    = request.form['ten_ct'].strip(),
            phong_ban = request.form['phong_ban'],
            ten_file  = request.form['ten_file'].strip(),
            duong_dan = request.form['duong_dan'].strip(),
            ngay_phat = datetime.strptime(request.form['ngay_phat'],'%Y-%m-%d').date(),
            thoi_luong= request.form.get('thoi_luong','').strip(),
            tai_tro   = request.form.get('tai_tro','Không tài trợ'),
            ghi_chu   = request.form.get('ghi_chu','').strip(),
            trang_thai= 'cho_duyet',
            nguoi_nop_id = session['user_id']
        )
        db.session.add(p)
        db.session.commit()
        # Auto-link to matching LichItem
        _auto_link_phieu(p)
        flash(f'Đã nộp phiếu: {p.ten_ct}', 'success')
        return redirect(url_for('phieu_list'))
    return render_template('phieu_form.html', phieu=None, default_phong=session.get('phong_ban',''))

def _auto_link_phieu(p):
    """Try to match phieu to a LichItem on same day."""
    items = LichItem.query.join(LichHDPS).filter(
        LichHDPS.ngay_phat == p.ngay_phat,
        LichItem.trang_thai == 'cho'
    ).all()
    name_clean = p.ten_ct.lower().strip()
    for item in items:
        if name_clean in item.ten_ct.lower() or item.ten_ct.lower() in name_clean:
            item.phieu_nt_id = p.id
            item.ten_file    = p.ten_file
            item.duong_dan   = p.duong_dan
            item.trang_thai  = 'co_file'
            db.session.commit()
            break

@app.route('/phieu/<int:pid>')
@login_required
def phieu_detail(pid):
    p = PhieuNghiemThu.query.get_or_404(pid)
    # Tìm LichItem liên kết
    linked = LichItem.query.filter_by(phieu_nt_id=pid).all()
    return render_template('phieu_detail.html', p=p, linked=linked)

@app.route('/phieu/<int:pid>/edit', methods=['GET','POST'])
@login_required
def phieu_edit(pid):
    p = PhieuNghiemThu.query.get_or_404(pid)
    if session['role'] not in ('admin','phongct') and \
       (p.nguoi_nop_id != session['user_id'] or p.trang_thai == 'da_duyet'):
        flash('Không thể sửa phiếu này.', 'danger')
        return redirect(url_for('phieu_detail', pid=pid))
    if request.method == 'POST':
        p.ten_ct    = request.form['ten_ct'].strip()
        p.phong_ban = request.form['phong_ban']
        p.ten_file  = request.form['ten_file'].strip()
        p.duong_dan = request.form['duong_dan'].strip()
        p.ngay_phat = datetime.strptime(request.form['ngay_phat'],'%Y-%m-%d').date()
        p.thoi_luong= request.form.get('thoi_luong','').strip()
        p.tai_tro   = request.form.get('tai_tro','Không tài trợ')
        p.ghi_chu   = request.form.get('ghi_chu','').strip()
        p.trang_thai= 'cho_duyet'
        p.updated_at= datetime.utcnow()
        db.session.commit()
        _auto_link_phieu(p)
        flash('Đã cập nhật phiếu.', 'success')
        return redirect(url_for('phieu_detail', pid=pid))
    return render_template('phieu_form.html', phieu=p, default_phong=p.phong_ban)

@app.route('/phieu/<int:pid>/duyet', methods=['POST'])
@login_required
@roles('admin','phongct')
def phieu_duyet(pid):
    p = PhieuNghiemThu.query.get_or_404(pid)
    p.trang_thai = 'da_duyet'; p.nguoi_duyet_id = session['user_id']
    p.updated_at = datetime.utcnow()
    # Update linked LichItem
    for item in LichItem.query.filter_by(phieu_nt_id=pid).all():
        if item.trang_thai == 'co_file': item.trang_thai = 'da_xac_nhan'
    db.session.commit()
    flash(f'Đã duyệt: {p.ten_ct}', 'success')
    return redirect(request.referrer or url_for('phieu_list'))

@app.route('/phieu/<int:pid>/tra_ve', methods=['POST'])
@login_required
@roles('admin','phongct')
def phieu_tra_ve(pid):
    p = PhieuNghiemThu.query.get_or_404(pid)
    p.trang_thai = 'tra_ve'; p.ly_do_tra = request.form.get('ly_do','')
    p.nguoi_duyet_id = session['user_id']; p.updated_at = datetime.utcnow()
    for item in LichItem.query.filter_by(phieu_nt_id=pid).all():
        item.trang_thai = 'cho'; item.phieu_nt_id = None
    db.session.commit()
    flash('Đã trả về phiếu.', 'warning')
    return redirect(request.referrer or url_for('phieu_list'))

# API: danh sách phiếu dạng JSON cho autocomplete
@app.route('/api/phieu_suggest')
@login_required
def api_phieu_suggest():
    ngay = request.args.get('ngay')
    q = PhieuNghiemThu.query.filter_by(trang_thai='da_duyet')
    if ngay:
        q = q.filter(PhieuNghiemThu.ngay_phat == datetime.strptime(ngay,'%Y-%m-%d').date())
    return jsonify([{'id':p.id,'ten':p.ten_ct,'file':p.ten_file,'path':p.duong_dan} for p in q.all()])

# ═══════════════════════════════════════════════════
#  COPY FILE
# ═══════════════════════════════════════════════════

@app.route('/copy')
@login_required
@roles('admin','phongct','phatsong')
def copy_list():
    jobs = CopyJob.query.order_by(CopyJob.created_at.desc()).limit(20).all()
    lichs = LichHDPS.query.filter(
        LichHDPS.trang_thai.in_(['hoan_chinh','da_xuat'])
    ).order_by(LichHDPS.ngay_phat.desc()).limit(14).all()
    default_dest = get_cfg('dest_default', r'V:\Thanhpham_PS_K1')
    return render_template('copy_list.html', jobs=jobs, lichs=lichs, default_dest=default_dest)

@app.route('/copy/start', methods=['POST'])
@login_required
@roles('admin','phongct','phatsong')
def copy_start():
    lid       = int(request.form['lich_id'])
    dest_path = request.form.get('dest_path', get_cfg('dest_default', r'V:\Thanhpham_PS_K1')).strip()
    set_cfg('dest_default', dest_path)
    lich = LichHDPS.query.get_or_404(lid)

    job = CopyJob(lich_id=lid, dest_path=dest_path,
                  trang_thai='cho', nguoi_tao_id=session['user_id'])
    # Count files
    items  = LichItem.query.filter_by(lich_id=lid).all()
    total  = sum(len([l for l in (i.ten_file or '').split('\n') if l.strip()])
                 for i in items if i.loai == 'file' and i.ten_file)
    job.tong_file = total
    db.session.add(job); db.session.commit()

    # Run in background thread (simulated on server without actual file access)
    t = threading.Thread(target=_run_copy, args=(job.id,), daemon=True)
    t.start()

    flash(f'Đã bắt đầu copy {total} file về {dest_path}', 'success')
    return redirect(url_for('copy_list'))

def _run_copy(job_id):
    """Background copy simulation — replace shutil.copy2 with actual network copy."""
    with app.app_context():
        job  = CopyJob.query.get(job_id)
        job.trang_thai = 'dang_chay'; db.session.commit()
        lich = job.lich
        items = LichItem.query.filter_by(lich_id=lich.id).all()
        log_lines = [f'Bắt đầu copy lúc {datetime.now().strftime("%H:%M:%S")}']
        copied = 0
        for item in items:
            if item.loai != 'file' or not item.ten_file: continue
            for fname in [f.strip() for f in item.ten_file.split('\n') if f.strip()]:
                if not fname.lower().endswith('.mpg'): fname += '.mpg'
                src = os.path.join(item.duong_dan or '', fname).replace('/', '\\')
                dst = os.path.join(job.dest_path, fname).replace('/', '\\')
                try:
                    # On actual Windows server: shutil.copy2(src, dst)
                    log_lines.append(f'[OK] {fname}')
                    copied += 1
                except Exception as e:
                    log_lines.append(f'[LỖI] {fname}: {e}')
                job.da_copy  = copied
                job.log      = '\n'.join(log_lines)
                db.session.commit()
                ttime.sleep(0.05)  # simulate IO
        job.trang_thai = 'hoan_thanh'
        log_lines.append(f'Hoàn thành {copied}/{job.tong_file} file lúc {datetime.now().strftime("%H:%M:%S")}')
        job.log = '\n'.join(log_lines); db.session.commit()

@app.route('/api/copy/<int:jid>/status')
@login_required
def copy_status(jid):
    j = CopyJob.query.get_or_404(jid)
    return jsonify(trang_thai=j.trang_thai, da_copy=j.da_copy,
                   tong_file=j.tong_file, log=j.log or '')

# ═══════════════════════════════════════════════════
#  USERS & SETTINGS
# ═══════════════════════════════════════════════════

@app.route('/users')
@login_required
@roles('admin')
def user_list():
    users = User.query.order_by(User.phong_ban, User.full_name).all()
    return render_template('user_list.html', users=users)

@app.route('/users/new', methods=['GET','POST'])
@login_required
@roles('admin')
def user_new():
    if request.method == 'POST':
        if User.query.filter_by(username=request.form['username']).first():
            flash('Tên đăng nhập đã tồn tại.', 'danger')
        else:
            u = User(username=request.form['username'].strip(),
                     full_name=request.form['full_name'].strip(),
                     phong_ban=request.form['phong_ban'],
                     role=request.form['role'])
            u.set_password(request.form['password'])
            db.session.add(u); db.session.commit()
            flash(f'Đã tạo tài khoản {u.full_name}.', 'success')
            return redirect(url_for('user_list'))
    return render_template('user_form.html', user=None)

@app.route('/users/<int:uid>/edit', methods=['GET','POST'])
@login_required
@roles('admin')
def user_edit(uid):
    u = User.query.get_or_404(uid)
    if request.method == 'POST':
        u.full_name = request.form['full_name'].strip()
        u.phong_ban = request.form['phong_ban']
        u.role      = request.form['role']
        u.active    = 'active' in request.form
        if request.form.get('password'): u.set_password(request.form['password'])
        db.session.commit(); flash('Đã cập nhật.', 'success')
        return redirect(url_for('user_list'))
    return render_template('user_form.html', user=u)

@app.route('/settings', methods=['GET','POST'])
@login_required
@roles('admin')
def settings():
    if request.method == 'POST':
        set_cfg('dest_default', request.form.get('dest_default',''))
        set_cfg('server_share', request.form.get('server_share',''))
        flash('Đã lưu cài đặt.', 'success')
    return render_template('settings.html',
        dest_default=get_cfg('dest_default', r'V:\Thanhpham_PS_K1'),
        server_share=get_cfg('server_share', r'\\server-22t01'))

# ═══════════════════════════════════════════════════
#  INIT DB
# ═══════════════════════════════════════════════════

def _migrate_columns():
    # Add new columns to existing tables (safe: skips if already exists)
    migrations = [
        ('lich_item', 'nguoi_thao_tac', 'VARCHAR(100)'),
        ('lich_item', 'file_size_mb', 'FLOAT'),
        ('lich_item', 'thoi_luong_giay', 'INTEGER'),
    ]
    with db.engine.connect() as conn:
        for table, col, typ in migrations:
            try:
                conn.execute(db.text(f'ALTER TABLE {table} ADD COLUMN {col} {typ}'))
                conn.commit()
            except Exception:
                pass  # Column already exists

def init_db():
    with app.app_context():
        db.create_all()
        _migrate_columns()
        if not User.query.filter_by(username='admin').first():
            demo = [
                ('admin',     'Quản trị hệ thống',     'Ban Giám đốc',        'admin',    'admin123'),
                ('bichloan',  'Lương Thị Bích Loan',    'Phòng CT Truyền hình','phongct',  'loan123'),
                ('maihuynh',  'Nguyễn Lê Mai Huỳnh',    'Phòng SXCT',          'phongcm',  'huynh123'),
                ('thanhsang', 'Huỳnh Thanh Sang',        'Phòng Biên dịch',     'phongcm',  'sang123'),
                ('vantuan',   'Trần Văn Thuận',          'Phòng Thời sự',       'phongcm',  'tuan123'),
                ('phatsong',  'Bộ phận Phát sóng',       'Kỹ thuật Phát sóng',  'phatsong', 'ps123'),
            ]
            for un,fn,pb,role,pw in demo:
                u = User(username=un, full_name=fn, phong_ban=pb, role=role)
                u.set_password(pw); db.session.add(u)
            db.session.commit()
            print('✓ Database khởi tạo xong.')

# DB duoc khoi tao lan dau khi co request (lazy init)
_db_initialized = False

if __name__ == '__main__':
    port = int(os.environ.get('PORT', 5000))
    app.run(host='0.0.0.0', port=port, debug=False)
